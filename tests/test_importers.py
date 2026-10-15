"""Behavioral tests for the human-reviewed import preview."""

from io import BytesIO
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from backend.importers import extract_candidates


def docx_bytes(build):
    document = Document()
    build(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def xlsx_bytes(build):
    workbook = Workbook()
    build(workbook)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


class ImporterTests(unittest.TestCase):
    def test_docx_keeps_unicode_and_document_order(self):
        def build(document):
            document.add_paragraph("  Saudação à tradição  ")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Ogum"
            table.cell(0, 1).text = "Oxóssi"
            document.add_paragraph("  Final  ")

        preview = extract_candidates(docx_bytes(build), "nomes.docx")
        self.assertEqual([item["text"] for item in preview["items"]],
                         ["Saudação à tradição", "Ogum", "Oxóssi", "Final"])
        self.assertEqual(preview["sheets"], [])
        self.assertEqual(preview["columns"], [])

    def test_xlsx_defaults_to_first_sheet_first_column_and_exposes_zero_based_columns(self):
        def build(workbook):
            sheet = workbook.active
            sheet.title = "Primeira"
            sheet.append(["Item", "Detalhe"])
            sheet.append(["Água", "Fonte"])
            sheet.append(["Vela", "Azul"])
            other = workbook.create_sheet("Outra")
            other.append(["Nome", "Quantidade"])
            other.append(["Axé", 3])

        content = xlsx_bytes(build)
        preview = extract_candidates(content, "itens.xlsx")
        self.assertEqual(preview["sheets"], ["Primeira", "Outra"])
        self.assertEqual(preview["columns"], [{"index": 0, "label": "Item"}, {"index": 1, "label": "Detalhe"}])
        self.assertEqual([item["text"] for item in preview["items"]], ["Item", "Água", "Vela"])
        selected = extract_candidates(content, "itens.xlsx", sheet="Outra", columns=[1, 0])
        self.assertEqual(selected["columns"], [{"index": 0, "label": "Nome"}, {"index": 1, "label": "Quantidade"}])
        self.assertEqual([item["text"] for item in selected["items"]], ["Nome — Quantidade", "Axé — 3"])

    def test_xlsx_does_not_emit_formula_without_cached_value(self):
        def build(workbook):
            sheet = workbook.active
            sheet.append(["Item"])
            sheet.append(["=1+1"])
            sheet.append(["real"])

        preview = extract_candidates(xlsx_bytes(build), "formula.xlsx")
        self.assertEqual([item["text"] for item in preview["items"]], ["Item", "real"])

    def test_pdf_without_extractable_text_has_helpful_error(self):
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        output = BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(ValueError, "digitalizado|texto"):
            extract_candidates(output.getvalue(), "scan.pdf")

    def test_text_pdf_yields_lines_for_review(self):
        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
        })
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 40 170 Td (Primeiro item) Tj 0 -20 Td (Segundo item) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = BytesIO()
        writer.write(output)
        preview = extract_candidates(output.getvalue(), "lista.pdf")
        self.assertEqual([item["text"] for item in preview["items"]],
                         ["Primeiro item", "Segundo item"])

    def test_xlsx_rejects_excessive_blank_rows(self):
        def build(workbook):
            sheet = workbook.active
            sheet.cell(row=1, column=1, value="item")
            sheet.cell(row=100002, column=1, value="outro")

        with self.assertRaisesRegex(ValueError, "linhas"):
            extract_candidates(xlsx_bytes(build), "linhas.xlsx")

    def test_pdf_over_100_pages_is_rejected(self):
        writer = PdfWriter()
        for _ in range(101):
            writer.add_blank_page(width=200, height=200)
        output = BytesIO()
        writer.write(output)
        with self.assertRaisesRegex(ValueError, "100 páginas"):
            extract_candidates(output.getvalue(), "livro.pdf")

    def test_empty_docx_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "vazio|texto"):
            extract_candidates(docx_bytes(lambda document: None), "vazio.docx")

    def test_legacy_and_unsupported_extensions_have_distinct_errors(self):
        for name in ("arquivo.doc", "arquivo.xls"):
            with self.subTest(name=name), self.assertRaisesRegex(ValueError, "Converta"):
                extract_candidates(b"old format", name)
        with self.assertRaisesRegex(ValueError, "não é aceito"):
            extract_candidates(b"data", "arquivo.exe")

    def test_mismatched_type_and_corrupt_office_file_are_rejected(self):
        valid = docx_bytes(lambda document: document.add_paragraph("Texto"))
        with self.assertRaisesRegex(ValueError, "formato|extensão"):
            extract_candidates(valid, "errado.xlsx")
        with self.assertRaisesRegex(ValueError, "corrompido|inválido"):
            extract_candidates(b"PK-not-a-zip", "falha.docx")

    def test_input_and_expanded_zip_limits_are_enforced(self):
        with self.assertRaisesRegex(ValueError, "10 MB"):
            extract_candidates(b"x" * (10 * 1024 * 1024 + 1), "enorme.pdf")
        output = BytesIO()
        with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
            archive.writestr("large.xml", b"x" * (50 * 1024 * 1024 + 1))
        with self.assertRaisesRegex(ValueError, "50 MB"):
            extract_candidates(output.getvalue(), "bomba.docx")

    def test_candidate_limit_rejects_excess_instead_of_truncating_silently(self):
        content = docx_bytes(lambda document: [document.add_paragraph(str(i)) for i in range(1001)])
        with self.assertRaisesRegex(ValueError, "1000"):
            extract_candidates(content, "muitos.docx")

    def test_invalid_sheet_or_column_is_rejected(self):
        content = xlsx_bytes(lambda workbook: workbook.active.append(["Única"]))
        with self.assertRaisesRegex(ValueError, "aba"):
            extract_candidates(content, "item.xlsx", sheet="Inexistente")
        with self.assertRaisesRegex(ValueError, "coluna"):
            extract_candidates(content, "item.xlsx", columns=[2])


if __name__ == "__main__":
    unittest.main()
