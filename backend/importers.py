"""Bounded, read-only extraction for the list import preview."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.table import Table
from openpyxl import load_workbook
from pypdf import PdfReader


MAX_INPUT = 10 * 1024 * 1024
MAX_EXPANDED_OFFICE = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_ITEMS = 1000
MAX_PDF_PAGES = 100
MAX_COLUMNS = 256
MAX_ROWS_SCANNED = 100_000
REVIEW_WARNING = "Revise os itens extraídos antes de criar a lista."


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def _add(items: list[dict[str, str]], value: object) -> None:
    text = _clean(value)
    if text:
        if len(items) >= MAX_ITEMS:
            raise ValueError("O arquivo tem mais de 1000 itens. Divida-o em arquivos menores.")
        items.append({"text": text})


def _check_office_archive(content: bytes, kind: str) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise ValueError("O arquivo Office contém entradas demais.")
            total = 0
            for entry in entries:
                if entry.flag_bits & 1:
                    raise ValueError("O arquivo está protegido por senha. Envie uma cópia sem senha.")
                total += entry.file_size
                if total > MAX_EXPANDED_OFFICE:
                    raise ValueError("O arquivo Office expandido ultrapassa 50 MB.")
            names = {entry.filename for entry in entries}
            required = "word/document.xml" if kind == ".docx" else "xl/workbook.xml"
            alternate = "xl/workbook.xml" if kind == ".docx" else "word/document.xml"
            if required not in names:
                if alternate in names:
                    raise ValueError("A extensão não corresponde ao formato do arquivo.")
                raise ValueError("O arquivo Office está corrompido ou é inválido.")
            if "[Content_Types].xml" not in names:
                raise ValueError("O arquivo Office está corrompido ou é inválido.")
            for entry in entries:
                if entry.filename.lower().endswith((".xml", ".rels")):
                    data = archive.read(entry)
                    if re.search(rb"<!\s*(?:DOCTYPE|ENTITY)\b", data, re.IGNORECASE):
                        raise ValueError("O arquivo Office contém XML inválido ou inseguro.")
    except BadZipFile as exc:
        raise ValueError("O arquivo Office está corrompido ou é inválido.") from exc


def _docx(content: bytes) -> dict:
    items: list[dict[str, str]] = []
    try:
        document = Document(BytesIO(content))
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            _add(items, paragraph.text)
            else:
                _add(items, block.text)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("O DOCX está corrompido ou é inválido.") from exc
    if not items:
        raise ValueError("O DOCX está vazio ou não contém texto para importar.")
    return {"items": items, "sheets": [], "columns": [], "warnings": [REVIEW_WARNING]}


def _xlsx(content: bytes, sheet: str | None, columns: list[int] | None) -> dict:
    workbook = None
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        sheet_names = workbook.sheetnames
        if not sheet_names:
            raise ValueError("A planilha está vazia.")
        if sheet is not None and sheet not in sheet_names:
            raise ValueError("A aba escolhida não existe na planilha.")
        worksheet = workbook[sheet or sheet_names[0]]
        if (worksheet.max_row or 0) > MAX_ROWS_SCANNED:
            raise ValueError("A planilha tem mais de 100 mil linhas. Divida-a em arquivos menores.")
        width = worksheet.max_column or 0
        if width > MAX_COLUMNS:
            raise ValueError("A planilha tem mais de 256 colunas.")
        if columns is None:
            chosen = [0]
        else:
            if not columns or any(type(index) is not int or index < 0 or index >= width for index in columns):
                raise ValueError("Escolha uma coluna válida da planilha (índices a partir de zero).")
            chosen = sorted(set(columns))
        iterator = worksheet.iter_rows(min_row=1, max_col=width, values_only=True)
        first_row = next(iterator, None)
        if first_row is None or width == 0:
            raise ValueError("A planilha está vazia.")
        available = [{"index": index, "label": _clean(first_row[index]) or f"Coluna {index + 1}"}
                     for index in range(width)]
        items: list[dict[str, str]] = []

        def add_row(row: tuple) -> None:
            values = [_clean(row[index]) for index in chosen]
            _add(items, " — ".join(value for value in values if value))

        add_row(first_row)
        for row_number, row in enumerate(iterator, start=2):
            if row_number > MAX_ROWS_SCANNED:
                raise ValueError("A planilha tem mais de 100 mil linhas. Divida-a em arquivos menores.")
            add_row(row)
        if not items:
            raise ValueError("A coluna escolhida está vazia ou não contém valores em cache.")
        return {"items": items, "sheets": sheet_names, "columns": available,
                "warnings": [REVIEW_WARNING, "Cabeçalhos podem ter sido incluídos; revise a primeira linha."]}
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("O XLSX está corrompido ou é inválido.") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _pdf(content: bytes) -> dict:
    try:
        reader = PdfReader(BytesIO(content), strict=True)
        if reader.is_encrypted:
            raise ValueError("O PDF está protegido por senha. Envie uma cópia sem senha.")
        if len(reader.pages) > MAX_PDF_PAGES:
            raise ValueError("O PDF ultrapassa o limite de 100 páginas.")
        items: list[dict[str, str]] = []
        for page in reader.pages:
            for line in (page.extract_text() or "").splitlines():
                _add(items, line)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("O PDF está corrompido ou é inválido.") from exc
    if not items:
        raise ValueError("O PDF não contém texto extraível. Se for digitalizado, converta-o com OCR antes de importar.")
    return {"items": items, "sheets": [], "columns": [], "warnings": [REVIEW_WARNING]}


def extract_candidates(content: bytes, filename: str, sheet: str | None = None,
                       columns: list[int] | None = None) -> dict:
    """Return candidate text for review; never persist or create a list.

    XLSX column indices are zero-based. The first row is included in candidates
    because the importer cannot know whether it is a header.
    """
    extension = Path(filename).suffix.lower()
    if extension in {".doc", ".xls"}:
        raise ValueError("Formato antigo não aceito. Converta o arquivo para DOCX ou XLSX.")
    if extension not in {".docx", ".xlsx", ".pdf"}:
        raise ValueError("Este tipo de arquivo não é aceito. Envie DOCX, XLSX ou PDF.")
    if not isinstance(content, bytes):
        raise ValueError("O conteúdo do arquivo é inválido.")
    if len(content) > MAX_INPUT:
        raise ValueError("O arquivo ultrapassa o limite de 10 MB.")
    if not content:
        raise ValueError("O arquivo está vazio.")
    if extension == ".pdf":
        if not content.startswith(b"%PDF-"):
            raise ValueError("A extensão não corresponde ao formato do arquivo PDF.")
        return _pdf(content)
    _check_office_archive(content, extension)
    if extension == ".docx":
        return _docx(content)
    return _xlsx(content, sheet, columns)
